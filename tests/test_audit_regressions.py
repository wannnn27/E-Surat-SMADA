"""Regression checks for operational bugs identified in the system audit."""

from __future__ import annotations

import csv
import io
import json
import sqlite3
import shutil
import unittest
import zipfile
from concurrent.futures import ThreadPoolExecutor
from contextlib import closing
from pathlib import Path
from unittest.mock import MagicMock, patch

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE

from esurat import application, database
from esurat.letters import _validate_request
from flask import session, request
from esurat.errors import RequestValidationError
from esurat.template_management import validate_custom_template
from tests import test_tu_workflow as workflow


class AuditRegressionTests(unittest.TestCase):
    setUp = workflow.TUWorkflowTests.setUp
    login = workflow.TUWorkflowTests.login
    form = workflow.TUWorkflowTests.form
    generate = workflow.TUWorkflowTests.generate

    def test_non_ascii_csrf_is_a_controlled_rejection(self):
        self.login()
        response = self.generate(token="token-\u00e9")
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.get_json()["code"], "csrf_invalid")

    def test_json_arrays_and_scalars_do_not_crash_routes(self):
        token = self.client.get("/api/csrf").get_json()["csrf_token"]
        for payload in ([1], "text", 12, None):
            response = self.client.post("/login", data=json.dumps(payload), content_type="application/json", headers={"X-CSRFToken": token})
            self.assertEqual(response.status_code, 400)
        token = self.login("admin-tu")
        for payload in ([1], "text", 12):
            response = self.client.post("/api/history/1/cancel", json=payload, headers={"X-CSRFToken": token})
            self.assertEqual(response.status_code, 400)

    def test_master_changes_cannot_change_contents_under_existing_number(self):
        self.login()
        form = self.form()
        first = self.generate(form=form)
        self.assertEqual(first.status_code, 200)
        person = self.app.extensions["esurat_data"]["guru_by_nip"][form["id_value"]]
        person["nama"] = "IDENTITAS BERUBAH"
        retry = self.generate(form=form)
        self.assertEqual(retry.status_code, 409)
        self.assertEqual(self.client.get("/api/list/riwayat").get_json()["total"], 1)

    def test_changed_password_invalidates_loaded_account_session(self):
        self.login()
        self.app.extensions["auth_users"]["operator-satu"]["password_hash"] = "changed-hash"
        self.assertEqual(self.client.get("/api/list/guru").status_code, 401)

    def test_explicit_missing_signer_cannot_silently_fall_back(self):
        from esurat.master_data import validate_master_data
        from esurat.errors import DataValidationError
        state = self.app.extensions["esurat_data"]
        with self.assertRaises(DataValidationError):
            validate_master_data(state["guru"], state["murid"], state["kode_arsip"], "999999999999999999")

    def test_duplicate_disabled_username_is_still_rejected(self):
        from esurat.security import _load_auth_users
        from esurat.errors import DataValidationError
        from pathlib import Path
        users = Path(self.temp.name) / "duplicate-users.json"
        users.write_text(json.dumps([{"username": "duplicate", "password_hash": "synthetic", "active": active} for active in (False, True)]), encoding="utf-8")
        with self.assertRaises(DataValidationError):
            _load_auth_users({"AUTH_USERS_FILE": str(users)})

    def test_old_fingerprint_is_rejected_without_allocating_a_number(self):
        self.login()
        form = self.form()
        self.assertEqual(self.generate(form=form).status_code, 200)
        with closing(sqlite3.connect(self.database)) as conn:
            conn.execute("UPDATE riwayat_surat SET payload_hash = ?", ("a" * 64,))
            conn.commit()
        self.assertEqual(self.generate(form=form).status_code, 409)
        self.assertEqual(self.client.get("/api/list/riwayat").get_json()["total"], 1)

    def prepare_builtin_copy(self):
        template_dir = Path(self.temp.name) / "builtin-templates"
        shutil.copytree(self.app.config["TEMPLATE_DIR"], template_dir)
        self.app.config["TEMPLATE_DIR"] = template_dir
        return template_dir / str(self.app.extensions["letter_registry"]["surat_keterangan_guru"]["template"])

    def test_live_builtin_change_is_rejected_before_number_allocation(self):
        self.login()
        path = self.prepare_builtin_copy()
        with zipfile.ZipFile(path, "a") as archive:
            archive.comment = b"new template revision"
        self.assertEqual(self.generate().status_code, 409)
        self.assertEqual(self.client.get("/api/list/riwayat").get_json()["total"], 0)

    def test_render_uses_verified_snapshot_if_template_changes_after_reservation(self):
        self.login()
        path = self.prepare_builtin_copy()
        original = application._render_letter
        def replace_after_reservation(validated, number):
            path.write_bytes(b"synthetic-invalid-template")
            return original(validated, number)
        with patch.object(application, "_render_letter", side_effect=replace_after_reservation):
            self.assertEqual(self.generate().status_code, 200)

    def test_cancellation_during_repeat_render_blocks_the_download(self):
        self.login()
        form = self.form()
        self.assertEqual(self.generate(form=form).status_code, 200)
        admin = self.app.test_client()
        token = self.login("admin-tu", admin)
        original = application._render_letter

        def cancel_during_render(validated, number):
            result = original(validated, number)
            response = admin.post("/api/history/1/cancel", json={"reason": "Surat salah tujuan"}, headers={"X-CSRFToken": token})
            self.assertEqual(response.status_code, 200)
            return result

        with patch.object(application, "_render_letter", side_effect=cancel_during_render):
            self.assertEqual(self.generate(form=form).status_code, 409)
        record = self.client.get("/api/list/riwayat").get_json()["items"][0]
        self.assertEqual(record["status"], "cancelled")
        with self.app.app_context():
            database._mark_letter_status(1, "failed", "late worker error")
            with self.assertRaises(RequestValidationError):
                database._mark_letter_status(1, "generated")
        self.assertEqual(self.client.get("/api/list/riwayat").get_json()["items"][0]["status"], "cancelled")

    def test_parallel_identical_requests_never_allocate_two_numbers(self):
        clients = [self.app.test_client(), self.app.test_client()]
        tokens = [self.login(client=client) for client in clients]
        form = self.form()
        with ThreadPoolExecutor(max_workers=2) as pool:
            responses = list(pool.map(lambda pair: self.generate(pair[0], form, pair[1]), zip(clients, tokens)))
        self.assertIn(200, [response.status_code for response in responses])
        self.assertTrue(all(response.status_code in {200, 409} for response in responses))
        self.login()
        self.assertEqual(self.client.get("/api/list/riwayat").get_json()["total"], 1)

    def test_postgres_reservation_locks_request_and_record(self):
        connection = MagicMock()
        connection.execute.return_value.fetchone.return_value = {"last_seq": 1, "id": 1}
        def execute(query, *args):
            cursor = MagicMock()
            cursor.fetchone.return_value = {"last_seq": 1, "id": 1} if ("RETURNING" in query) else None
            return cursor
        connection.execute.side_effect = execute
        with self.app.test_request_context("/generate", method="POST", data=self.form()):
            session.update(authenticated=True, username="operator-satu", role="operator")
            validated = _validate_request(request.form, preview=False)
            self.app.config["DATABASE"] = "postgresql://example.invalid/database"
            with patch.object(database, "_connect_db", return_value=connection):
                reservation = database._reserve_letter(validated)
        self.assertEqual(reservation["id"], 1)
        statements = [str(call.args[0]) for call in connection.execute.call_args_list]
        self.assertTrue(any("pg_advisory_xact_lock" in query for query in statements))
        self.assertTrue(any("request_id = %s FOR UPDATE" in query for query in statements))
        connection.commit.assert_called_once()

    def test_unique_constraint_conflict_returns_actionable_409(self):
        connection = MagicMock()
        connection.execute.side_effect = sqlite3.IntegrityError("unique constraint")
        with self.app.test_request_context("/generate", method="POST", data=self.form()):
            session.update(authenticated=True, username="operator-satu", role="operator")
            validated = _validate_request(request.form, preview=False)
            with patch.object(database, "_connect_db", return_value=connection):
                with self.assertRaises(RequestValidationError) as caught:
                    database._reserve_letter(validated)
        self.assertEqual(caught.exception.status_code, 409)
        connection.rollback.assert_called_once()

    def test_month_dashboard_uses_wib_boundaries_and_excludes_future_month(self):
        self.login()
        for _ in range(3):
            self.assertEqual(self.generate().status_code, 200)
        with closing(sqlite3.connect(self.database)) as conn:
            for record_id, timestamp in enumerate(("2026-07-31T18:00:00+00:00", "2026-07-31T16:59:59+00:00", "2026-08-31T17:00:00+00:00"), start=1):
                conn.execute("UPDATE riwayat_surat SET created_at = ? WHERE id = ?", (timestamp, record_id))
            conn.commit()
        html = self.client.get("/admin").get_data(as_text=True)
        marker = html.index("Aktivitas bulan ini")
        self.assertIn("<strong>1</strong>", html[marker:marker + 180])

    def test_csv_neutralizes_whitespace_prefixed_formulas(self):
        self.login()
        self.assertEqual(self.generate().status_code, 200)
        for unsafe in ("  =1+1", "\t=1+1", "\r=1+1", "@SUM(1)"):
            with closing(sqlite3.connect(self.database)) as conn:
                conn.execute("UPDATE riwayat_surat SET keperluan = ?", (unsafe,))
                conn.commit()
            data = self.client.get("/api/history/export.csv").get_data(as_text=True)
            rows = list(csv.reader(io.StringIO(data.lstrip("\ufeff"))))
            self.assertEqual(rows[1][7], "'" + unsafe)

    def template(self, external=False):
        document = Document()
        document.add_paragraph("{{ nomor_surat }} {{ tanggal_surat }} {{ nama }} {{ keperluan }}")
        if external:
            document.part.relate_to("https://example.invalid", RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    def validate_template(self, content):
        return validate_custom_template({"key": "audit_surat", "label": "Surat audit", "description": "Template pengujian sintetis", "category": "guru", "person_mode": "single", "default_code": "800.1.11", "signer": "kepsek"}, content, {"800.1.11"})

    def test_oversized_zip_rejected_before_decompression(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", "a" * (21 * 1024 * 1024))
        with patch.object(zipfile.ZipFile, "testzip", side_effect=AssertionError("must not decompress")):
            with self.assertRaises(RequestValidationError) as caught:
                self.validate_template(buffer.getvalue())
        self.assertIn("terlalu besar", caught.exception.field_errors["template_file"])

    def test_external_relationship_with_single_quotes_is_rejected(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(self.template(external=True))) as source, zipfile.ZipFile(buffer, "w") as target:
            for item in source.infolist():
                target.writestr(item, source.read(item.filename).replace(b'TargetMode="External"', b"TargetMode='External'"))
        with self.assertRaises(RequestValidationError) as caught:
            self.validate_template(buffer.getvalue())
        self.assertIn("eksternal", caught.exception.field_errors["template_file"])

    def test_macro_disguised_as_docx_is_rejected(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(self.template())) as source, zipfile.ZipFile(buffer, "w") as target:
            for item in source.infolist():
                target.writestr(item, source.read(item.filename))
            target.writestr("word/vbaProject.bin", b"synthetic-macro")
        with self.assertRaises(RequestValidationError) as caught:
            self.validate_template(buffer.getvalue())
        self.assertIn("macro", caught.exception.field_errors["template_file"])

    def test_custom_field_names_must_match_safe_form_contract(self):
        document = Document()
        document.add_paragraph("{{ nomor_surat }} {{ tanggal_surat }} {{ nama }} {{ unsupportedField }}")
        buffer = io.BytesIO()
        document.save(buffer)
        with self.assertRaises(RequestValidationError) as caught:
            self.validate_template(buffer.getvalue())
        self.assertIn("field tambahan", caught.exception.field_errors["template_file"])

    def test_duplicate_features_removed_and_icons_are_local(self):
        self.login()
        html = self.client.get("/").get_data(as_text=True)
        self.assertNotIn('id="modalRiwayat"', html)
        self.assertIn('href="/admin/history"', html)
        self.assertNotIn("cdnjs.cloudflare.com", html)
        self.assertIn("vendor/fontawesome/css/solid.min.css", html)
        self.assertNotIn("guru-staf", self.client.get("/admin/master-data").get_data(as_text=True))

    def test_disabled_template_leaves_source_and_metadata_in_database(self):
        token = self.login("admin-tu")
        metadata = {"key": "audit_surat", "label": "Surat audit", "description": "Template pengujian sintetis", "category": "guru", "person_mode": "single", "default_code": "800.1.11", "signer": "kepsek", "template_file": (io.BytesIO(self.template()), "audit.docx"), "_csrf_token": token}
        self.assertEqual(self.client.post("/admin/templates", data=metadata).status_code, 302)
        self.assertEqual(self.client.post("/admin/templates/audit_surat/delete", data={"confirm": "DELETE", "_csrf_token": token}).status_code, 302)
        self.assertEqual(self.client.get("/api/fields/audit_surat").status_code, 404)
        with closing(sqlite3.connect(self.database)) as conn:
            row = conn.execute("SELECT active, length(content) FROM custom_templates WHERE key = 'audit_surat'").fetchone()
        self.assertIsNotNone(row)
        self.assertEqual(row[0], 0)
        self.assertGreater(row[1], 0)

